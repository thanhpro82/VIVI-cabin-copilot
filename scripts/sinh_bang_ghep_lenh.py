"""Sinh `docs/kich_ban_test_ghep_lenh.xlsx` — lưới tổ hợp "nói nhiều lệnh một lượt".

Chạy:  .\\.venv\\Scripts\\python.exe scripts\\sinh_bang_ghep_lenh.py

Năm cột do MÁY điền (router thật + policy thật), bốn cột để NGƯỜI điền lúc ngồi test.
Cột "Kỳ vọng của NGƯỜI" viết tay trong chính file này và **cố ý không bằng** cột máy đo:
chỗ hai cột lệch nhau là chỗ đáng mở issue. Một bảng chỉ chép lại hành vi router thì
không tìm ra được gì.

Sinh lại sau mỗi đợt merge để biết ô nào đã đổi. Lệnh này dựng lại toàn bộ file, nhưng
**bốn cột người điền (K–N) được chép sang bản mới theo mã VT-P** — xem `doc_cot_nguoi`.
Trước 21/08 nó ghi đè trắng, và một buổi test tay là một buổi ngồi điền lại.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.agents.policy import classify
from src.agents.router import DeterministicControlRouter
from src.fixtures import load_poi_fixture

STOP = {"motion": {"speed_kph": 0, "gear": "P"}, "state_version": 1}
DRIVE = {"motion": {"speed_kph": 45, "gear": "D"}, "state_version": 2}

#: (nhóm, câu nói, số ý người nói, kỳ vọng của một người dùng bình thường).
#:
#: Cột cuối là chỗ **cố ý không hỏi router**. Nó ghi lại điều một người bình thường
#: nghĩ sẽ xảy ra sau khi nói câu đó — kể cả khi biết chắc hệ thống hiện không làm
#: được. Mất cột này thì bảng chỉ còn là ảnh chụp hành vi router.
CAU: list[tuple[str, str, int, str]] = [
    ("1. Hai ý cùng HVAC", "Bật điều hòa 22 độ và quạt gió mức 2", 2, "Bật máy + 22 độ + quạt mức 2"),
    ("1. Hai ý cùng HVAC", "Quạt gió mức 2 và điều hòa 22 độ", 2, "Như trên, chỉ đảo thứ tự"),
    ("1. Hai ý cùng HVAC", "Đặt quạt gió mức 1 rồi điều hòa 26 độ", 2, "Quạt mức 1 + 26 độ"),
    ("1. Hai ý cùng HVAC", "Điều hòa 24 độ với quạt gió mức 3", 2, "24 độ + quạt mức 3"),
    ("1. Hai ý cùng HVAC", "Điều hòa 24 độ cùng quạt gió mức 3", 2, "24 độ + quạt mức 3"),
    ("1. Hai ý cùng HVAC", "Tắt điều hòa và quạt gió", 2, "Tắt máy + quạt về 0"),
    ("1. Hai ý cùng HVAC", "Tắt điều hòa, tắt quạt gió", 2, "Tắt máy + quạt về 0 (không liên từ)"),
    ("1. Hai ý cùng HVAC", "Bật điều hòa 20 độ quạt gió mức 3", 2, "20 độ + quạt mức 3 (không liên từ)"),
    ("1. Hai ý cùng HVAC", "Bật điều hòa 22 độ và quạt gió mức 9", 2, "Từ chối ĐÚNG vế quạt; 22 độ vẫn nên chạy"),
    ("1. Hai ý cùng HVAC", "Bật điều hòa 45 độ và quạt gió mức 2", 2, "Từ chối ĐÚNG vế nhiệt độ; quạt vẫn nên chạy"),
    ("2. Chéo domain, 2 ý", "Bật điều hòa và mở nhạc", 2, "Bật điều hòa + phát nhạc"),
    ("2. Chéo domain, 2 ý", "Mở nhạc và bật điều hòa", 2, "Phát nhạc + bật điều hòa"),
    ("2. Chéo domain, 2 ý", "Bật điều hòa 24 độ và bật đèn trần", 2, "24 độ + đèn trần bật"),
    ("2. Chéo domain, 2 ý", "Bật đèn trần và bật điều hòa 24 độ", 2, "Đèn trần bật + 24 độ"),
    ("2. Chéo domain, 2 ý", "Phát nhạc và bật đèn chiếu gần", 2, "Nhạc chạy + đèn cốt bật"),
    ("2. Chéo domain, 2 ý", "Bật đèn chiếu gần và phát nhạc", 2, "Đèn cốt bật + nhạc chạy"),
    ("2. Chéo domain, 2 ý", "Bật điều hòa và mở cửa sổ bên lái 30 phần trăm", 2, "Điều hòa chạy ngay + hỏi duyệt kính"),
    ("2. Chéo domain, 2 ý", "Mở cửa sổ bên lái 30 phần trăm và bật điều hòa", 2, "Hỏi duyệt kính + điều hòa chạy"),
    ("2. Chéo domain, 2 ý", "Mở cửa bên lái và mở cốp sau", 2, "MỘT hộp thoại duyệt liệt kê cả cửa lẫn cốp"),
    ("2. Chéo domain, 2 ý", "Mở cốp sau và mở cửa bên lái", 2, "Như trên, đảo thứ tự"),
    ("2. Chéo domain, 2 ý", "Bật sưởi ghế lái mức 2 và bật điều hòa 22 độ", 2, "Sưởi mức 2 + điều hòa 22 độ"),
    ("2. Chéo domain, 2 ý", "Dẫn đường đến trạm sạc và phát nhạc", 2, "Đặt điểm đến + nhạc chạy"),
    ("2. Chéo domain, 2 ý", "Phát nhạc và dẫn đường đến trạm sạc", 2, "Nhạc chạy + đặt điểm đến"),
    ("2. Chéo domain, 2 ý", "Mở YouTube và bật điều hòa", 2, "Mở app + điều hòa bật"),
    ("2. Chéo domain, 2 ý", "Bật điều hòa và mở YouTube", 2, "Điều hòa bật + mở app"),
    ("2. Chéo domain, 2 ý", "Mở cửa sổ bên lái 30 phần trăm và mở cửa bên lái", 2, "Duyệt gộp: kính 30% + cửa lái"),
    ("2. Chéo domain, 2 ý", "Đặt âm lượng 40 và bật đèn trần", 2, "Âm lượng 40 + đèn trần"),
    ("2. Chéo domain, 2 ý", "Ngả lưng ghế lái 70 phần trăm và bật sưởi ghế lái mức 2", 2, "Duyệt ngả ghế + sưởi mức 2 chạy ngay"),
    ("3. Cùng domain, 2 tool", "Phát nhạc và đặt âm lượng 40", 2, "Nhạc chạy + âm lượng 40"),
    ("3. Cùng domain, 2 tool", "Đặt âm lượng 40 và chuyển bài", 2, "Âm lượng 40 + sang bài kế"),
    ("3. Cùng domain, 2 tool", "Bật đèn trần và bật đèn chiếu xa", 2, "Đèn trần + đèn pha xa"),
    ("3. Cùng domain, 2 tool", "Mở cửa sổ bên lái 30 phần trăm và cửa sổ bên phụ 50 phần trăm", 2, "Hai kính, hai mức, một lần duyệt"),
    ("3. Cùng domain, 2 tool", "Mở cửa bên lái và đóng cửa bên phụ", 2, "Mở một cửa, đóng cửa kia"),
    ("3. Cùng domain, 2 tool", "Bật sưởi ghế lái mức 2 và sưởi ghế phụ mức 1", 2, "Hai ghế, hai mức"),
    ("4. Ba ý trở lên", "Bật điều hòa 22 độ, quạt gió mức 2 và phát nhạc", 3, "22 độ + quạt 2 + nhạc"),
    ("4. Ba ý trở lên", "Bật điều hòa 22 độ và quạt gió mức 2 và mở nhạc và bật đèn trần", 4, "22 độ + quạt 2 + nhạc + đèn trần"),
    ("4. Ba ý trở lên", "Phát nhạc, bật đèn trần và mở cửa sổ bên lái 30 phần trăm", 3, "Nhạc + đèn + duyệt kính"),
    ("4. Ba ý trở lên", "Bật điều hòa, phát nhạc, bật đèn trần", 3, "Ba lệnh S1 chạy hết"),
    ("4. Ba ý trở lên", "Mở cửa bên lái, mở cốp sau và hạ kính bên lái 40 phần trăm", 3, "Một hộp thoại duyệt cho cả ba"),
    ("5. Lệnh + câu hỏi", "Bật điều hòa và áp suất lốp khuyến nghị là bao nhiêu", 2, "Điều hòa bật + trả lời câu hỏi"),
    ("5. Lệnh + câu hỏi", "Áp suất lốp khuyến nghị là bao nhiêu và bật điều hòa", 2, "Trả lời câu hỏi + điều hòa bật"),
    ("5. Lệnh + câu hỏi", "Bật đèn trần và đèn chào mừng là gì", 2, "Đèn trần bật + trả lời câu hỏi"),
    ("5. Lệnh + câu hỏi", "Cách sạc pin thế nào và mở cốp sau", 2, "Trả lời + hỏi duyệt mở cốp"),
    ("6. Lệnh + phủ định", "Bật điều hòa nhưng đừng mở cửa sổ", 2, "Điều hòa BẬT, kính không đụng tới"),
    ("6. Lệnh + phủ định", "Đừng mở cửa sổ, bật điều hòa", 2, "Kính không đụng, điều hòa BẬT"),
    ("6. Lệnh + phủ định", "Bật điều hòa và đừng phát nhạc", 2, "Điều hòa BẬT, nhạc không phát"),
    ("6. Lệnh + phủ định", "Đừng bật điều hòa và đừng phát nhạc", 2, "Không làm gì, nhưng nói rõ đã hiểu"),
    ("7. Trộn mức an toàn", "Bật điều hòa và mở cửa bên lái", 2, "S1 chạy ngay + hộp thoại duyệt cửa"),
    ("7. Trộn mức an toàn", "Mở cửa bên lái và bật điều hòa", 2, "Hộp thoại duyệt cửa + S1 chạy"),
    ("7. Trộn mức an toàn", "Mở cửa sổ bên lái 30 phần trăm và mở cốp sau", 2, "Một hộp thoại: kính + cốp"),
    ("7. Trộn mức an toàn", "Mở tất cả cửa và mở tất cả cửa sổ 50 phần trăm", 2, "Một hộp thoại: 4 cửa + 4 kính = 8 bước"),
    ("7. Trộn mức an toàn", "Mở YouTube và mở cửa bên lái", 2, "Mở app + hộp thoại duyệt cửa"),
    ("8. Thiếu slot một vế", "Bật điều hòa và mở cửa sổ", 2, "Điều hòa bật + hỏi lại kính nào"),
    ("8. Thiếu slot một vế", "Mở cửa sổ và bật điều hòa", 2, "Hỏi lại kính nào + điều hòa bật"),
    ("8. Thiếu slot một vế", "Phát nhạc và bật đèn", 2, "Nhạc chạy + hỏi đèn pha hay đèn trần"),
    ("8. Thiếu slot một vế", "Bật đèn và phát nhạc", 2, "Hỏi đèn nào + nhạc chạy"),
    ("8. Thiếu slot một vế", "Chỉnh điều hòa và phát nhạc", 2, "Hỏi bao nhiêu độ + nhạc chạy"),
    ("9. Liên từ lạ", "Bật điều hòa xong mở nhạc", 2, "Điều hòa + nhạc"),
    ("9. Liên từ lạ", "Bật điều hòa sau đó mở nhạc", 2, "Điều hòa + nhạc"),
    ("9. Liên từ lạ", "Bật điều hòa luôn tiện mở nhạc", 2, "Điều hòa + nhạc"),
    ("9. Liên từ lạ", "Bật điều hòa; mở nhạc", 2, "Điều hòa + nhạc"),
    ("9. Liên từ lạ", "Bật điều hòa mở nhạc", 2, "Điều hòa + nhạc (không liên từ)"),
    ("9. Liên từ lạ", "Vừa bật điều hòa vừa mở nhạc", 2, "Điều hòa + nhạc"),
    ("9. Liên từ lạ", "Bật điều hòa và cả mở nhạc nữa", 2, "Điều hòa + nhạc"),
    ("10. Lịch sự / dài dòng", "Làm ơn bật điều hòa và phát nhạc giúp tôi", 2, "Điều hòa + nhạc"),
    ("10. Lịch sự / dài dòng", "Bạn bật điều hòa 22 độ rồi mở nhạc lên nhé", 2, "22 độ + nhạc"),
    ("10. Lịch sự / dài dòng", "Cho tôi xin điều hòa 22 độ và nhạc", 2, "22 độ + nhạc"),
    ("11. Tương đối / mơ hồ", "Tăng điều hòa và tăng âm lượng", 2, "Hỏi lại CẢ HAI con số"),
    ("11. Tương đối / mơ hồ", "Bật điều hòa 22 độ và tăng âm lượng thêm 10", 2, "22 độ chạy + hỏi lại âm lượng tuyệt đối"),
    ("11. Tương đối / mơ hồ", "Bật điều hòa và mở nó ra", 2, "Điều hòa bật + hỏi lại 'nó' là gì"),
    ("12. Lặp cùng một ý", "Bật điều hòa và bật điều hòa", 1, "Bật một lần, không chạy hai lần"),
    ("12. Lặp cùng một ý", "Đặt điều hòa 22 độ và đặt điều hòa 26 độ", 2, "Hỏi lại: hai con số mâu thuẫn"),
    ("12. Lặp cùng một ý", "Mở cửa sổ bên lái 30 phần trăm và mở cửa sổ bên lái 70 phần trăm", 2, "Hỏi lại: cùng kính, hai mức"),
    ("13. Trong + ngoài phạm vi", "Bật điều hòa và bật cruise control", 2, "Điều hòa bật + nói rõ KHÔNG làm được cruise control"),
    ("13. Trong + ngoài phạm vi", "Bật cruise control và bật điều hòa", 2, "Nói rõ không làm được + điều hòa bật"),
    ("13. Trong + ngoài phạm vi", "Phát nhạc và gọi cho vợ tôi", 2, "Nhạc chạy + nói rõ không gọi điện được"),
    ("13. Trong + ngoài phạm vi", "Bật điều hòa và tắt túi khí", 2, "Điều hòa bật + TỪ CHỐI tắt túi khí"),
    ("13. Trong + ngoài phạm vi", "Tắt túi khí và bật điều hòa", 2, "TỪ CHỐI túi khí + điều hòa bật"),
]

COT: list[tuple[str, int]] = [
    ("Mã", 9),
    ("Nhóm tổ hợp", 22),
    ("Câu nói", 56),
    ("Số ý người nói", 8),
    ("Kỳ vọng của NGƯỜI (viết tay)", 44),
    ("Router đo được", 13),
    ("Lý do (reason)", 25),
    ("Bước sinh ra", 38),
    ("An toàn: đứng yên → đang chạy", 15),
    ("Ý bị bỏ im lặng?", 12),
    ("NGHE / THẤY THẬT (bạn tự điền)", 40),
    ("Đạt?", 9),
    ("trace_id", 20),
    ("Ghi chú", 32),
]

DAM = Font(bold=True, color="FFFFFF")
NEN_TIEU_DE = PatternFill("solid", fgColor="1F3864")
NEN_MAY = PatternFill("solid", fgColor="E8EDF5")
NEN_NGUOI = PatternFill("solid", fgColor="FFF7E6")
NEN_NUOT = PatternFill("solid", fgColor="FBD5D5")
NEN_MAT = PatternFill("solid", fgColor="F8CBAD")

COT_MAY = (6, 7, 8, 9, 10)
COT_NGUOI = (5, 11, 12, 13, 14)


def doc_cot_nguoi(dich: Path) -> dict[str, list[Any]]:
    """Đọc bốn cột người điền (K–N) của bản cũ, khoá theo mã `VT-P**`.

    Không có hàm này thì mỗi lần sinh lại là xoá trắng công của người vừa ngồi test —
    và hệ quả thật là **không ai dám sinh lại**, nên bảng trôi khỏi hành vi router
    đúng vào lúc nó cần đúng nhất.

    Khoá theo **mã**, không theo số dòng: thêm một câu vào `CAU` là mọi dòng phía dưới
    xê dịch, và chép theo dòng sẽ gán kết quả test của câu này sang câu khác — sai còn
    tệ hơn mất trắng, vì nó trông vẫn hợp lý.
    """
    if not dich.exists():
        return {}
    try:
        cu = load_workbook(dich)["Ghép lệnh"]
    except (KeyError, OSError, ValueError):
        # File hỏng/khác định dạng thì sinh mới, đừng chặn cả lệnh vì bản cũ.
        return {}
    giu: dict[str, list[Any]] = {}
    for hang in cu.iter_rows(min_row=2, values_only=True):
        ma = hang[0]
        if not ma:
            continue
        nguoi = [hang[j - 1] if j - 1 < len(hang) else None for j in (11, 12, 13, 14)]
        if any(gia not in (None, "") for gia in nguoi):
            giu[str(ma)] = nguoi
    return giu


def dem_viec(buoc: tuple[Any, ...]) -> int:
    """Số **việc khác nhau** trong kế hoạch — đếm theo tool, không theo bước.

    Đếm bước là chỗ bảng cũ nói dối. `"Mở tất cả cửa và mở tất cả cửa sổ 50 phần trăm"`
    ra 4 bước `set_window_position` cho **một** ý, nên `4 >= 2` và cột J chấm "không bỏ
    ý" — trong khi cả bốn cửa đã biến mất. Một ý nở ra nhiều bước luôn dùng **cùng một**
    tool, nên đếm tool khác nhau là bịt đúng lỗ đó.

    Vẫn không phải trọng tài, và không thể là: `"bật điều hòa 22 độ"` là **một** ý mà ra
    hai tool (bật máy + đặt nhiệt). Cột E — kỳ vọng viết tay — mới là trọng tài.
    """
    return len({b.tool for b in buoc})


def main() -> None:
    dich = Path(__file__).resolve().parents[1] / "docs" / "kich_ban_test_ghep_lenh.xlsx"
    cot_nguoi = doc_cot_nguoi(dich)

    # `poi_fixture` **phải** truyền vào: thiếu nó thì `_poi_ids` rỗng, guard
    # `poi_not_in_fixture` không chạy lần nào, và mọi hàng điều hướng trong bảng mô tả
    # một router khác với router đang chạy thật. Đúng lỗi mà issue #171 đã đóng ở
    # `graph.py:305` — cùng một cái bẫy, chỉ khác chỗ đặt.
    router = DeterministicControlRouter(poi_fixture=load_poi_fixture())
    wb = Workbook()
    ws = wb.active
    ws.title = "Ghép lệnh"

    for i, (ten, rong) in enumerate(COT, start=1):
        o = ws.cell(row=1, column=i, value=ten)
        o.font = DAM
        o.fill = NEN_TIEU_DE
        o.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = rong
    ws.row_dimensions[1].height = 34

    dong = 2
    for i, (nhom, cau, so_y, ky_vong) in enumerate(CAU, start=1):
        quyet_dinh = router.route(cau)
        buoc = quyet_dinh.candidate_plan.steps if quyet_dinh.candidate_plan else ()

        muc = "—"
        if buoc:
            dung_yen = "/".join(sorted({classify(s.tool, STOP) for s in buoc}))
            dang_chay = "/".join(sorted({classify(s.tool, DRIVE) for s in buoc}))
            muc = dung_yen if dung_yen == dang_chay else f"{dung_yen} → {dang_chay}"

        chi_tiet = "\n".join(f"{s.tool} {s.args}" for s in buoc) or "—"

        # Phán đoán CƠ HỌC, không phải trọng tài: cột "Kỳ vọng của NGƯỜI" mới là
        # trọng tài. Một lượt "số bước khớp số ý" vẫn có thể sai nội dung.
        if so_y >= 2 and quyet_dinh.disposition == "control" and dem_viec(buoc) < so_y:
            nuot = "CÓ"
        elif so_y >= 2 and quyet_dinh.disposition in {"not_control", "denied", "clarify"}:
            nuot = "CẢ LƯỢT"
        else:
            nuot = "không"

        ma = f"VT-P{i:02d}"
        da_dien = cot_nguoi.get(ma, ["", "", "", ""])
        gia_tri = [
            ma,
            nhom,
            cau,
            so_y,
            ky_vong,
            quyet_dinh.disposition,
            quyet_dinh.reason or "—",
            chi_tiet,
            muc,
            nuot,
            *[gia if gia is not None else "" for gia in da_dien],
        ]
        for j, gia in enumerate(gia_tri, start=1):
            o = ws.cell(row=dong, column=j, value=gia)
            o.alignment = Alignment(vertical="top", wrap_text=j in (3, 5, 8, 11, 14))
            if j in COT_MAY:
                o.fill = NEN_MAY
            elif j in COT_NGUOI:
                o.fill = NEN_NGUOI
        if nuot == "CÓ":
            ws.cell(row=dong, column=10).fill = NEN_NUOT
        elif nuot == "CẢ LƯỢT":
            ws.cell(row=dong, column=10).fill = NEN_MAT
        dong += 1

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:N{dong - 1}"
    chon = DataValidation(type="list", formula1='"Đạt,Trượt,Chưa chạy"', allow_blank=True)
    ws.add_data_validation(chon)
    chon.add(f"L2:L{dong - 1}")

    huong_dan = wb.create_sheet("Cách dùng")
    huong_dan.column_dimensions["A"].width = 28
    huong_dan.column_dimensions["B"].width = 110
    noi_dung = [
        ("Bảng này để làm gì", "Tìm chỗ HỎNG của lượt ghép nhiều lệnh — không phải để xác nhận lại những gì đã biết là chạy được."),
        ("", ""),
        ("Cột xanh (F–J)", "MÁY điền. Sinh bằng scripts/sinh_bang_ghep_lenh.py từ router + policy THẬT, không phải kỳ vọng."),
        ("Cột vàng (E, K–N)", "NGƯỜI điền. Cột E viết trước khi chạy; K–N bạn tự điền lúc ngồi test."),
        ("", ""),
        ("Đọc cột J thế nào", "CÓ = router chạy ÍT VIỆC hơn số ý người nói → phần còn lại biến mất và xe KHÔNG nói gì về nó."),
        ("", "CẢ LƯỢT = không ý nào chạy: cả câu rơi xuống sổ tay, bị từ chối, hoặc bị hỏi lại."),
        ("", "không = số việc khớp số ý. KHÔNG có nghĩa là đúng — vẫn phải nghe câu trả lời rồi mới kết luận."),
        ("", ""),
        ("Cách test một hàng", "1) Đọc cột E TRƯỚC, đừng đọc cột F — biết trước router làm gì thì mắt sẽ tự bào chữa cho nó."),
        ("", "2) Nói câu ở cột C vào IVI (hoặc gõ, nếu đang đo router chứ không đo STT)."),
        ("", "3) Ghi vào cột K đúng thứ NGHE và THẤY: xe nói gì, màn đổi gì, xe có thực sự làm không."),
        ("", "4) Đạt/Trượt so với cột E, KHÔNG so với cột F. 'Đúng theo thiết kế' mà người dùng mất một lệnh vẫn là Trượt."),
        ("", "5) Trượt thì lấy trace_id ở màn /engineer bỏ vào cột M. Trace là bằng chứng; ảnh chụp màn hình thì không."),
        ("", ""),
        ("Vì sao E ≠ F mới là điểm chính", "Chỗ hai cột lệch nhau là chỗ đáng mở issue. Bảng nào chỉ chép lại hành vi router thì không tìm ra được gì."),
        ("", ""),
        ("Sinh lại sau khi merge", ".\\.venv\\Scripts\\python.exe scripts\\sinh_bang_ghep_lenh.py"),
        ("", "Bốn cột K–N bạn đã điền được CHÉP SANG bản mới theo mã VT-P, nên sinh lại không mất công test."),
        ("", "Cột A–J thì dựng lại hoàn toàn: đó là ảnh chụp router hiện tại, không phải dữ liệu của bạn."),
        ("Nguồn", "docs/kich_ban_test_thoai.md §9b. Mức an toàn: docs/safety_and_hitl.md. Bản đồ điều khiển được: docs/coverage_matrix.md."),
    ]
    for i, (trai, phai) in enumerate(noi_dung, start=1):
        o = huong_dan.cell(row=i, column=1, value=trai)
        o.font = Font(bold=True)
        o.alignment = Alignment(vertical="top")
        huong_dan.cell(row=i, column=2, value=phai).alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(dich)

    tong = dong - 2
    nuot_co = sum(1 for r in range(2, dong) if ws.cell(row=r, column=10).value == "CÓ")
    mat = sum(1 for r in range(2, dong) if ws.cell(row=r, column=10).value == "CẢ LƯỢT")
    print(f"da ghi: {dich}")
    print(f"tong {tong} cau | bo im lang {nuot_co} | mat ca luot {mat} | so viec khop {tong - nuot_co - mat}")
    print(f"giu lai {len(cot_nguoi)} hang da dien tay (cot K-N)")


if __name__ == "__main__":
    main()
